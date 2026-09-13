from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).parents[1]
OUT = ROOT / "output/docx/InterpretaAI-Proposta-MVP.docx"
BLUE, DARK, YELLOW, PALE, GREEN, INK, RED = "2E74B5", "1F4D78", "FFD21E", "EAF2F8", "E8F5E9", "172033", "FCE8E6"

def shade(cell, color):
    node = OxmlElement("w:shd"); node.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(node)

def cell_margin(cell, value=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar"); tc_pr.append(tc_mar)
    for tag in ("top", "start", "bottom", "end"):
        node = OxmlElement("w:" + tag); node.set(qn("w:w"), str(value)); node.set(qn("w:type"), "dxa")
        tc_mar.append(node)

def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("INTERPRETAAI  •  MVP  |  ")
    run.font.size = Pt(9); run.font.color.rgb = RGBColor.from_string(DARK)
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); paragraph._p.append(field)

def heading(text, level=1, kicker=None):
    if kicker:
        p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(2)
        r = p.add_run(kicker.upper()); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string(DARK)
    return doc.add_heading(text, level=level)

def body(text):
    return doc.add_paragraph(text)

def bullets(items):
    for item in items: doc.add_paragraph(item, style="List Bullet")

def callout(title, text, color=PALE):
    table = doc.add_table(rows=1, cols=1); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
    cell = table.cell(0, 0); shade(cell, color); cell_margin(cell, 140)
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(title.upper() + "\n"); r.bold = True; r.font.color.rgb = RGBColor.from_string(DARK)
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def status_table(rows):
    table = doc.add_table(rows=1, cols=2); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
    table.columns[0].width, table.columns[1].width = Cm(4.1), Cm(12.3)
    for i, value in enumerate(("ESTADO", "EVIDÊNCIA / LIMITE")):
        cell = table.cell(0, i); cell.text = value; shade(cell, BLUE); cell_margin(cell)
        for run in cell.paragraphs[0].runs: run.bold = True; run.font.color.rgb = RGBColor(255, 255, 255)
    for label, description, color in rows:
        cells = table.add_row().cells; cells[0].text = label; cells[1].text = description; shade(cells[0], color)
        for cell in cells:
            cell_margin(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cant_split = OxmlElement("w:cantSplit"); cell._tc.get_or_add_tcPr().append(cant_split)
        cells[0].paragraphs[0].runs[0].bold = True
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    return table

def image_pair(left, right, cap_left, cap_right, height=Inches(2.55)):
    table = doc.add_table(rows=2, cols=2); table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, path in enumerate((left, right)):
        p = table.cell(0, i).paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), height=height)
    for i, caption in enumerate((cap_left, cap_right)):
        p = table.cell(1, i).paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(caption); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string(DARK)

def page():
    doc.add_page_break()

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(2.2)
section.header_distance = section.footer_distance = Cm(.9)
section.different_first_page_header_footer = True

normal = doc.styles["Normal"]
normal.font.name = "Calibri"; normal.font.size = Pt(11); normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.space_after = Pt(8); normal.paragraph_format.line_spacing = 1.333
for name, size, color, before, after in (
    ("Title", 30, DARK, 0, 12), ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 12, 6), ("Heading 3", 12, DARK, 8, 4)):
    style = doc.styles[name]; style.font.name = "Calibri"; style.font.size = Pt(size)
    style.font.bold = True; style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
for name in ("List Bullet", "List Number"):
    style = doc.styles[name]; style.font.name = "Calibri"; style.font.size = Pt(11)
    style.paragraph_format.left_indent = Inches(.375); style.paragraph_format.first_line_indent = Inches(-.194)
    style.paragraph_format.space_after = Pt(4); style.paragraph_format.line_spacing = 1.208

header = section.header.paragraphs[0]; header.text = "INTERPRETAAI  /  PROPOSTA DE MVP"
header.runs[0].bold = True; header.runs[0].font.size = Pt(9); header.runs[0].font.color.rgb = RGBColor.from_string(DARK)
page_number(section.footer.paragraphs[0]); page_number(section.first_page_footer.paragraphs[0])

# 1 — capa editorial
bar = doc.add_table(rows=1, cols=1); bar.autofit = False; bar.columns[0].width = Cm(16.6)
shade(bar.cell(0,0), BLUE); cell_margin(bar.cell(0,0), 230)
r = bar.cell(0,0).paragraphs[0].add_run("INTERPRETA AI"); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = RGBColor(255,255,255)
doc.add_paragraph().paragraph_format.space_after = Pt(32)
p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.add_run("Alfabetização que\nconversa com a criança")
p = doc.add_paragraph(); r = p.add_run("Proposta técnica e pedagógica do MVP")
r.bold = True; r.font.size = Pt(16); r.font.color.rgb = RGBColor.from_string(BLUE); p.paragraph_format.space_after = Pt(28)
callout("Método LEIA", "Ler • Escrever • Interpretar • Aplicar", YELLOW)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.space_before = Pt(34)
p.add_run("COAUTORIA GUIADA\n").bold = True
p.add_run("A criança ajuda a LEIA e os personagens a concluir a história enquanto desenvolve linguagem, interpretação e participação.")
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(38)
r = p.add_run("VERSÃO MVP 0.1  •  13 DE SETEMBRO DE 2026"); r.bold = True; r.font.color.rgb = RGBColor.from_string(DARK)

# 2
page(); heading("Problema e proposta de valor", kicker="01 • Por que existe")
body("A alfabetização não acontece apenas quando a criança acerta uma letra. Ela se fortalece quando a criança observa, nomeia, compara, formula hipóteses e aplica o que percebeu em uma situação conhecida. Em sala, o professor precisa conduzir grupos, acolher ritmos diferentes e ainda registrar sinais de aprendizagem. Uma interface baseada em texto, rolagem e respostas binárias amplia a distância para quem ainda não lê.")
callout("Pergunta central", "Como o produto ajuda a otimizar o fluxo da alfabetização sem substituir o professor nem transformar a criança em usuária de um chatbot aberto?")
callout("Aderência ao desafio", "O smartphone deixa de competir pela atenção: participa de um ciclo curto de aprendizagem, colaboração e criatividade e depois descansa para a atividade continuar no mundo real.", YELLOW)
heading("Resposta do InterpretaAI", 2)
body("O InterpretaAI transforma a criança em ajudante da LEIA. A história apresenta uma situação familiar — bola, escola, ônibus ou alimento — e pede uma contribuição curta por voz ou toque. A mediação reconhece o esforço, devolve uma pergunta e encaminha a próxima etapa. O professor recebe observações de participação, modalidade, ajuda e tempo; não recebe um veredito automático.")
bullets([
    "Menos explicação operacional: voz e destaque visual indicam o próximo passo.",
    "Menos abandono: uma decisão por tela e reconexão suave após inatividade.",
    "Mais tempo pedagógico: registros locais resumem o percurso para o professor.",
    "Mais inclusão: voz, toque, contraste e redução de estímulos coexistem.",
])
body("O valor do MVP não está em gerar conteúdo ilimitado. Está em provar um ciclo curto, repetível e observável: a criança entra, compreende o convite, participa da história, relaciona som e palavra, conclui uma tarefa e deixa sinais úteis para a intervenção docente.")
callout("Recorte do MVP", "Uma história central, cinco cenas expressivas e três figuras de quebra-cabeça. Poucas possibilidades, mas todas demonstráveis e coerentes.", GREEN)

# 3
page(); heading("Método LEIA", kicker="02 • Estrutura pedagógica")
body("LEIA é o coração do produto e organiza tanto a experiência da criança quanto a leitura das métricas. Cada atividade atravessa quatro movimentos visíveis, ainda que a criança não conheça seus nomes escritos.")
table = doc.add_table(rows=1, cols=3); table.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, value in enumerate(("ETAPA", "NO PRODUTO", "SINAL PEDAGÓGICO")):
    cell=table.cell(0,i); cell.text=value; shade(cell,BLUE); cell_margin(cell)
    for run in cell.paragraphs[0].runs: run.bold=True; run.font.color.rgb=RGBColor(255,255,255)
for row in (
    ("Ler", "Observar a cena e ouvir o diálogo.", "Atenção a personagens, objetos e sequência."),
    ("Escrever", "Montar BOLA com letras faladas.", "Relação entre grafema, nome e fonema."),
    ("Interpretar", "Explicar o que pode ajudar Lia.", "Hipótese contextual e expressão oral."),
    ("Aplicar", "Representar a cena com o grupo.", "Transferência, escuta e cooperação.")):
    cells=table.add_row().cells
    for i,value in enumerate(row): cells[i].text=value; cell_margin(cells[i])
    shade(cells[0],YELLOW)
heading("Exemplo completo",2)
body("Lia não encontra a bola. A criança primeiro vê a imagem e ouve Lia e Davi. Depois pode dizer sua ideia à LEIA ou tocar em um caminho preparado. A reação não declara uma emoção como certa: mostra como procurar junto ou perguntar o que aconteceu podem fazer a narrativa avançar. Em seguida, a palavra BOLA é construída letra a letra e a turma representa a solução.")
callout("Regra de mediação", "Reconhecer a contribuição → conectar com a cena → fazer somente uma pergunta seguinte. Nunca diagnosticar, dar nota, usar culpa ou reduzir interpretação a certo/errado.", GREEN)
body("Essa organização permite trocar o tema sem trocar a arquitetura. Frutas, água, chuva, carro ou futebol tornam-se contextos de histórias; o ciclo pedagógico permanece estável e mensurável.")

# 4
page(); heading("Experiência da criança e dinâmica em grupo", kicker="03 • Jornada")
image_pair(ROOT/"output/screenshots/gibi-observar-360x640.png", ROOT/"output/screenshots/gibi-conversa-360x640.png", "1. Observar e ouvir", "2. Contar uma ideia", Inches(2.65))
body("O gibi foi separado em três telas por cena. Na primeira, a ilustração e os balões ocupam o centro; a ação principal é confirmar que observou. Na segunda, a pergunta aparece com microfone e duas alternativas visuais. Na terceira, a LEIA reage e oferece o avanço. A criança nunca precisa procurar o botão abaixo da dobra.")
heading("Como entra na sala de aula",2)
bullets([
    "Duplas: uma criança observa e a outra conta a ideia; depois trocam os papéis.",
    "Pequenos grupos: cada grupo escolhe um caminho e explica à turma.",
    "Rotação: um grupo usa o tablet enquanto outros montam palavras ou representam.",
    "Fechamento: o professor compara estratégias sem eleger emoção única ou ranking.",
])
body("O quebra-cabeça 2×2 ou 3×2 introduz uma pausa motora e visual. Ao concluir, a criança ouve e pronuncia a palavra. Brincar deixa de ser um prêmio desconectado e integra o percurso entre imagem, fala, escrita e uso social.")
body("No fechamento, a LEIA avisa que o celular já ajudou e pode descansar na mesa. A dupla continua a conversa, a representação ou a busca de objetos sem a tela, tornando o uso consciente uma ação observável e não apenas um discurso.")
callout("Inovação central", "Coautoria guiada: a criança sente que ajuda a LEIA e os personagens; a mediação adapta a próxima pergunta sem retirar do professor a condução pedagógica.", YELLOW)

# 5
page(); heading("Design infantil: atenção sem sobrecarga", kicker="04 • Interface e neurodiversidade")
image_pair(ROOT/"output/screenshots/before/home-360x640.png", ROOT/"output/screenshots/home-360x640.png", "ANTES • CTA dependia de rolagem", "DEPOIS • ação principal visível", Inches(2.65))
body("A auditoria mostrou o problema da versão anterior: em 360×640, o acesso ao Modo Escola ficava fora do viewport e “VER MAIS” competia com a tarefa. O novo ChildStageScaffold usa BoxWithConstraints, reduz ilustração e espaçamento antes do texto e mantém fonte mínima de 16sp e alvos de toque de 48dp.")
heading("Padrões aplicados",2)
bullets([
    "Uma decisão principal por estado; conteúdo extenso vira etapa controlada.",
    "Botão pressionado comprime, clareia e produz retorno sonoro curto.",
    "Quatro sons originais em SoundPool: toque, troca, descoberta e celebração.",
    "Fonemas continuam falados; bipes não substituem informação pedagógica.",
    "Reduzir estímulos remove sons e partículas, preservando voz, contraste e direção.",
])
body("A reconexão respeita autorregulação. Aos 20 segundos, a indicação visual reaparece; aos 40, uma única fala convida: “Ei, detetive! A história está esperando a sua ideia. Vamos juntos?”. O relógio pausa durante escuta, resposta e segundo plano e reinicia com interação. Não há culpa, cobrança ou repetição insistente.")
callout("Bem-estar digital", "Foco, reconexão sem culpa, estímulos reduzidos e descanso da tela são proteções pedagógicas e socioemocionais. O MVP não oferece tratamento, avaliação psicológica ou diagnóstico.", GREEN)

# 6
page(); heading("IA e arquitetura técnica", kicker="05 • Componentes e fluxo")
arch=doc.add_table(rows=1,cols=5); arch.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,(label,color) in enumerate((("Android\nCompose",YELLOW),("HTTPS\n6 s",PALE),("Spring Boot\nJava 17",BLUE),("Ollama\nQwen 2.5",GREEN),("Kokoro\nWAV pt-BR",GREEN))):
    cell=arch.cell(0,i); cell.text=label; shade(cell,color); cell_margin(cell,80); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in cell.paragraphs[0].runs:
        run.bold=True
        if color==BLUE: run.font.color.rgb=RGBColor(255,255,255)
body("Fluxo: o Android reconhece a fala e envia sessionId, sceneId, turno, transcrição, personagem e preferência de estímulos. O servidor valida o contrato, recupera no máximo seis mensagens por dez minutos, chama a mediação e sintetiza a fala. A resposta contém texto, personagem, áudio Base64, reação visual, próxima ação, categoria de observação e indicador degraded.")
heading("Arquitetura do MVP",2)
bullets([
    "Android Kotlin/Compose: interface, SpeechRecognizer, visão local, áudio temporário e métricas SQLite.",
    "Spring Boot 3.5.16 e Java 17: endpoint versionado, validação e logs sem conteúdo infantil.",
    "LangChain4j 1.20.0 + Ollama: Qwen 2.5 3B local, temperatura 0,2 e resposta JSON.",
    "Kokoro: vozes pt-BR feminina e masculina executadas no microservidor, sem custo por chamada.",
])
callout("Limite deliberado", "Sem agentes, RAG, banco vetorial ou LangGraph. Para o MVP, previsibilidade e testabilidade valem mais que autonomia ampla.")
status_table([
    ("IMPLEMENTADO", "Qwen e Kokoro respondem localmente nas duas vozes; fallback preserva o contrato.", GREEN),
    ("DEMONSTRADO", "Quick Tunnel HTTPS validado e APK configurado para o endereço público temporário.", PALE),
    ("PENDENTE", "Servidor estável na Oracle, autenticação, limite de requisições e monitoramento.", RED),
])

# 7
page(); heading("Sistema de vozes e fala relacional", kicker="06 • Presença da LEIA")
body("A voz não é apenas leitura de tela. Ela exerce um papel relacional com começo, continuidade e respeito ao silêncio. A voz principal é feminina e pertence à LEIA e à narradora. O microservidor também produz uma voz masculina para Davi; o roteamento completo de todos os diálogos por personagem ainda será validado no roteiro.")
table=doc.add_table(rows=1,cols=3); table.alignment=WD_TABLE_ALIGNMENT.CENTER
for i,value in enumerate(("TIPO","QUANDO","EXEMPLO")):
    cell=table.cell(0,i); cell.text=value; shade(cell,BLUE); cell_margin(cell)
    for run in cell.paragraphs[0].runs: run.bold=True; run.font.color.rgb=RGBColor(255,255,255)
for row in (
    ("Chamada inicial","Uma vez ao entrar","Oi! Eu sou a LEIA. Quer me ajudar a descobrir o que aconteceu?"),
    ("Interação","Após a ideia","Gostei da sua ideia! O que mais você percebe nessa cena?"),
    ("Reconexão","Uma vez, aos 40 s","Ei, detetive! A história está esperando a sua ideia. Vamos juntos?")):
    cells=table.add_row().cells
    for i,value in enumerate(row): cells[i].text=value; cell_margin(cells[i])
    shade(cells[0],YELLOW)
heading("Personagens e fallback",2)
bullets([
    "LEIA_FEMALE: Kokoro pf_dora, voz principal feminina.",
    "DAVI_MALE: Kokoro pm_alex, diálogos masculinos.",
    "Sem resposta em seis segundos: “A LEIA está sem internet, mas continua com você.”",
])
body("O áudio WAV é temporário: o Android grava no cache apenas para reprodução e apaga ao terminar. Nenhuma credencial viaja no APK. A síntese principal roda no Mac com Kokoro; se o servidor não responder, o TTS instalado no Android mantém a atividade compreensível e identifica o modo degradado.")
callout("Critério de segurança", "Respostas com no máximo duas frases, uma pergunta seguinte e três interações por sessão. O prompt proíbe nota, diagnóstico, culpa e classificação absoluta de emoção.", GREEN)

# 8
page(); heading("Professor e secretaria", kicker="07 • Métricas que apoiam decisão")
body("A experiência infantil é simples; a área adulta pode ser rolável e densa. O painel local já mostra sessões, respostas, modalidade de voz, observações no gibi, ciclos LEIA, quebra-cabeças, pedidos de ajuda e tempo médio. Esses sinais apoiam planejamento e conversa pedagógica — não produzem laudo, nota automática ou ranking.")
heading("Leitura em três níveis",2)
status_table([
    ("CRIANÇA", "Percurso no dispositivo: participação, ajuda, modalidade e duração.", YELLOW),
    ("PROFESSOR", "Visão por atividade e turma para formar grupos e escolher intervenções.", PALE),
    ("SECRETARIA", "Agregados por escola/turma/professor; sem áudio bruto e sem ranking infantil.", GREEN),
])
heading("Uso em grupos",2)
body("O professor pode formar grupos por necessidade observável: crianças que se beneficiam de mais apoio fonêmico, crianças que formulam hipóteses oralmente mas evitam escrita, ou grupos que precisam praticar escuta e turnos. O dado não determina o grupo sozinho; ele sugere onde observar. A secretaria acompanha adesão, conclusão e demanda de apoio para distribuir formação e recursos.")
bullets([
    "Fluxo: início, conclusão, abandono por etapa e tempo de resposta.",
    "Pedagogia: modalidade, ajuda solicitada, interpretação registrada e ciclo aplicado.",
    "Implantação: tablets ativos, turmas participantes e qualidade de sincronização.",
])
callout("Fronteira verificável", "Sincronização autenticada e dashboards da secretaria são evolução futura. O MVP mantém métricas de participação no tablet, sem substituir observação docente por precisão, nota ou ranking.")

# 9
page(); heading("Privacidade, segurança e Modo Foco", kicker="08 • Uso responsável")
body("O produto trata dados de crianças e deve operar por minimização. O APK registra eventos pedagógicos locais sem coluna de acerto, não salva áudio bruto e limita a transcrição a 280 caracteres. O servidor registra duração, turno e fallback, nunca transcrição ou áudio. A memória fica em RAM, retém seis mensagens e descarta sessões antigas na próxima atividade do servidor.")
heading("Modo Foco: o que é real",2)
body("No emulador provisionado como Device Owner, a auditoria confirmou mLockTaskModeState=LOCKED. Nesse modo, o pacote é allowlisted, Home e Recentes ficam bloqueados, a atividade se torna Home persistente e o sistema pode aplicar Não Perturbe após autorização. Em aparelho comum, Android permite apenas fixação de tela mediante confirmação adulta; um APK sozinho não obtém silenciosamente privilégios de Device Owner.")
bullets([
    "RECORD_AUDIO e CAMERA são solicitados no momento do uso e com explicação adulta.",
    "ACCESS_NOTIFICATION_POLICY exige concessão explícita na configuração do Android.",
    "O APK contém apenas a URL HTTPS; Ollama e Kokoro não ficam expostos diretamente.",
    "Área do educador usa PIN de demonstração; produção exige identidade institucional.",
    "SpeechRecognizer prefere operação offline, mas o mecanismo do Android/OEM pode usar rede.",
])
callout("Limite do MVP", "PIN fixo, endpoint sem autenticação e limpeza de cache/memória ainda dependente do ciclo de execução impedem uso com dados reais. São adequados à demonstração, não à implantação.", RED)
heading("Princípios para evolução",2)
body("Antes do piloto real: avaliação de impacto de privacidade, perfis de acesso, política de retenção, consentimento conforme contexto escolar, criptografia de sincronização, resposta a incidentes e validação com educadores e famílias. A IA deve continuar mediadora contextual, nunca avaliadora autônoma da criança.")

# 10
page(); heading("Estado do MVP, validação e próximos passos", kicker="09 • Fechamento")
status_table([
    ("ADEQUAÇÃO AO TEMA", "Foco, ciclo curto e celular em descanso para continuar em grupo.", GREEN),
    ("ORIGINALIDADE E INOVAÇÃO", "Coautoria por voz em gibi, fonema, puzzle e mundo real.", PALE),
    ("SOLUÇÃO TECNOLÓGICA", "APK, visão local, Spring/LangChain4j, Qwen/Kokoro e fallback.", YELLOW),
    ("UTILIDADE E APLICABILIDADE", "Uma decisão por tela e métricas de participação; piloto ainda necessário.", PALE),
])
heading("Riscos e próximos passos",2)
bullets([
    "Conteúdo estreito: prova um ciclo LEIA, não um currículo completo.",
    "Infraestrutura temporária: migrar para Oracle com autenticação, limites e monitoramento.",
    "Campo não validado: testar voz, câmera, ruído, acessibilidade e compreensão em turma pequena.",
    "Secretaria/família são futuras; bem-estar digital não é cuidado clínico; privacidade e consentimento exigem avaliação.",
])
heading("Conclusão",2)
body("O InterpretaAI transforma uma distração potencial em ferramenta breve de alfabetização e devolve a experiência ao grupo, com evidências reais e limites declarados.")
heading("Referências essenciais",2)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.line_spacing=1.0; p.paragraph_format.space_after=Pt(2)
for text in (
    "LangChain4j. Integração Ollama — docs.langchain4j.dev/integrations/language-models/ollama.",
    "Kokoro-82M. Síntese de voz open-weight — github.com/hexgrad/kokoro.",
    "Google ML Kit. Image labeling e text recognition on-device — developers.google.com/ml-kit.",
    "Android Developers. Lock task mode — developer.android.com/work/dpc/dedicated-devices/lock-task-mode.",
    "Repositório InterpretaAI. Testes, contrato API e evidências visuais, versão 0.1."):
    run=p.add_run(text+"\n"); run.font.size=Pt(8.5)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
